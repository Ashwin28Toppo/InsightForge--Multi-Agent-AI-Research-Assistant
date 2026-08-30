"""Unit tests for document ingestion (PDF/DOCX/TXT).

All fixtures are generated locally (temporary files); no external APIs or
third-party PDF files are used.
"""
from __future__ import annotations

import uuid

import pytest
from docx import Document as DocxDocument

from backend.app.rag.ingestion import (
    DocumentNotFoundError,
    UnsupportedExtensionError,
    ingest_document,
)
from backend.app.rag.schemas import DocumentContent


# ── Fixture helper: minimal PDF generator ────────────────────────────────────

def _escape_pdf_string(value: str) -> str:
    """Escape a string for use inside a PDF literal string."""
    return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _make_pdf(page_texts: list[str]) -> bytes:
    """Build a minimal valid PDF with one page per entry in ``page_texts``.

    Object layout:
      1 = catalog, 2 = pages tree,
      per page i: page obj (3 + 2i), content stream obj (4 + 2i),
      last obj = shared Helvetica font.
    """
    page_count = len(page_texts)
    font_number = 3 + 2 * page_count
    num_objects = font_number

    definitions: dict[int, bytes] = {}
    definitions[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(page_count))
    definitions[2] = f"<< /Type /Pages /Kids [{kids}] /Count {page_count} >>".encode("ascii")

    for index, text in enumerate(page_texts):
        page_number = 3 + 2 * index
        stream_number = 4 + 2 * index
        definitions[page_number] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_number} 0 R >> >> "
            f"/Contents {stream_number} 0 R >>"
        ).encode("ascii")
        content = (
            f"BT /F1 12 Tf 72 720 Td ({_escape_pdf_string(text)}) Tj ET"
        ).encode("latin-1")
        definitions[stream_number] = (
            b"<< /Length " + str(len(content)).encode("ascii") + b" >>\nstream\n"
            + content + b"\nendstream"
        )

    definitions[font_number] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"

    out = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for number in range(1, num_objects + 1):
        offsets[number] = len(out)
        out += f"{number} 0 obj\n".encode("ascii")
        out += definitions[number]
        out += b"\nendobj\n"

    xref_position = len(out)
    out += f"xref\n0 {num_objects + 1}\n".encode("ascii")
    out += b"0000000000 65535 f \n"
    for number in range(1, num_objects + 1):
        out += f"{offsets[number]:010d} 00000 n \n".encode("ascii")
    out += (
        f"trailer\n<< /Size {num_objects + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_position}\n%%EOF\n"
    ).encode("ascii")
    return bytes(out)


# ── TXT ingestion ────────────────────────────────────────────────────────────

def test_ingest_txt(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("  Hello   world.\n\n\n\nSecond paragraph.  ", encoding="utf-8")

    doc = ingest_document(path)

    assert isinstance(doc, DocumentContent)
    assert doc.source_type == "txt"
    assert doc.source_name == "notes.txt"
    assert len(doc.pages) == 1
    page = doc.pages[0]
    assert page.page_number is None
    assert "Hello world." in page.text
    assert "Second paragraph." in page.text
    assert "\n\n\n\n" not in page.text


def test_ingest_txt_unicode(tmp_path):
    path = tmp_path / "unicode.txt"
    path.write_text("Café — naïve résumé: 日本語テスト", encoding="utf-8")

    doc = ingest_document(path)

    assert "Café" in doc.pages[0].text
    assert "日本語テスト" in doc.pages[0].text


def test_ingest_txt_empty(tmp_path):
    path = tmp_path / "empty.txt"
    path.write_text("   \n\n  ", encoding="utf-8")

    doc = ingest_document(path)

    assert len(doc.pages) == 1
    assert doc.pages[0].text == ""


# ── DOCX ingestion ───────────────────────────────────────────────────────────

def test_ingest_docx(tmp_path):
    path = tmp_path / "report.docx"
    document = DocxDocument()
    document.add_paragraph("First paragraph")
    document.add_paragraph("Second paragraph")
    document.save(str(path))

    doc = ingest_document(path)

    assert doc.source_type == "docx"
    assert doc.source_name == "report.docx"
    assert len(doc.pages) == 1
    page = doc.pages[0]
    assert page.page_number is None
    # Paragraph ordering is preserved.
    assert page.text.index("First paragraph") < page.text.index("Second paragraph")


# ── PDF ingestion ────────────────────────────────────────────────────────────

def test_ingest_pdf_preserves_page_numbers(tmp_path):
    path = tmp_path / "sample.pdf"
    path.write_bytes(_make_pdf(["Hello from page one", "Hello from page two"]))

    doc = ingest_document(path)

    assert doc.source_type == "pdf"
    assert len(doc.pages) == 2
    assert doc.pages[0].page_number == 1
    assert doc.pages[1].page_number == 2
    assert "Hello from page one" in doc.pages[0].text
    assert "Hello from page two" in doc.pages[1].text


def test_ingest_pdf_empty_page_handled_gracefully(tmp_path):
    path = tmp_path / "blank.pdf"
    path.write_bytes(_make_pdf(["Has text", ""]))

    doc = ingest_document(path)

    assert len(doc.pages) == 2
    assert doc.pages[1].page_number == 2
    assert doc.pages[1].text == ""


# ── Validation & errors ──────────────────────────────────────────────────────

def test_ingest_document_id_is_uuid4(tmp_path):
    path = tmp_path / "a.txt"
    path.write_text("hello", encoding="utf-8")

    first = ingest_document(path)
    second = ingest_document(path)

    assert uuid.UUID(first.document_id)
    assert first.document_id != second.document_id


def test_ingest_unsupported_extension(tmp_path):
    path = tmp_path / "notes.md"
    path.write_text("# Title", encoding="utf-8")

    with pytest.raises(ValueError) as exc_info:
        ingest_document(path)
    assert "Unsupported" in str(exc_info.value)
    assert isinstance(exc_info.value, UnsupportedExtensionError)


def test_ingest_missing_file(tmp_path):
    path = tmp_path / "does_not_exist.txt"

    with pytest.raises(FileNotFoundError) as exc_info:
        ingest_document(path)
    assert isinstance(exc_info.value, DocumentNotFoundError)


def test_ingest_directory_rejected(tmp_path):
    with pytest.raises(FileNotFoundError):
        ingest_document(tmp_path)
